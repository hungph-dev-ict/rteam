import openpyxl
import os
import zipfile
import base64
import requests

def perform_image_ocr(image_bytes: bytes, filename: str, llm) -> str:
    """Sends image bytes directly to active LLM provider to perform OCR."""
    if not image_bytes or llm is None:
        return ""
        
    img_b64 = base64.b64encode(image_bytes).decode("utf-8")
    provider_type = llm.__class__.__name__.lower()
    
    # Determine mime type from filename
    mime_type = "image/png"
    if filename.lower().endswith(".jpg") or filename.lower().endswith(".jpeg"):
        mime_type = "image/jpeg"
    elif filename.lower().endswith(".gif"):
        mime_type = "image/gif"
        
    prompt = (
        "Extract all text from this image which is embedded in a candidate document. "
        "Preserve layout structure, columns, tables, headers, and bullet points. "
        "Write the exact text content found. Do not summarize and do not explain."
    )
    
    try:
        if "google" in provider_type or "gemini" in provider_type:
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{llm.model}:generateContent?key={llm.api_key}"
            payload = {
                "contents": [{
                    "parts": [
                        {"text": prompt},
                        {"inlineData": {"mimeType": mime_type, "data": img_b64}}
                    ]
                }],
                "generationConfig": {"temperature": 0.1}
            }
            res = requests.post(url, json=payload, headers={"Content-Type": "application/json"}, timeout=90)
            res.raise_for_status()
            data = res.json()
            return data["candidates"][0]["content"]["parts"][0]["text"].strip()
            
        elif "openai" in provider_type:
            model = llm.model
            if "gpt-3.5" in model:
                model = "gpt-4o-mini"
            url = "https://api.openai.com/v1/chat/completions"
            payload = {
                "model": model,
                "messages": [{
                    "role": "user",
                    "content": [
                        {"type": "text", "text": prompt},
                        {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{img_b64}"}}
                    ]
                }],
                "temperature": 0.1
            }
            headers = {"Authorization": f"Bearer {llm.api_key}", "Content-Type": "application/json"}
            res = requests.post(url, json=payload, headers=headers, timeout=90)
            res.raise_for_status()
            data = res.json()
            return data["choices"][0]["message"]["content"].strip()
            
        elif "claude" in provider_type or "anthropic" in provider_type:
            url = "https://api.anthropic.com/v1/messages"
            payload = {
                "model": llm.model,
                "max_tokens": 3000,
                "messages": [{
                    "role": "user",
                    "content": [
                        {
                            "type": "image",
                            "source": {"type": "base64", "media_type": mime_type, "data": img_b64}
                        },
                        {"type": "text", "text": prompt}
                    ]
                }],
                "temperature": 0.1
            }
            headers = {
                "x-api-key": llm.api_key,
                "anthropic-version": "2023-06-01",
                "content-type": "application/json"
            }
            res = requests.post(url, json=payload, headers=headers, timeout=90)
            res.raise_for_status()
            data = res.json()
            return data["content"][0]["text"].strip()
            
        elif "ollama" in provider_type:
            url = f"{llm.host}/api/chat"
            payload = {
                "model": llm.model,
                "messages": [{
                    "role": "user",
                    "content": prompt,
                    "images": [img_b64]
                }],
                "stream": False,
                "options": {"temperature": 0.1}
            }
            res = requests.post(url, json=payload, timeout=90)
            res.raise_for_status()
            data = res.json()
            return data["message"]["content"].strip()
            
    except Exception as e:
        return f"[Error executing AI OCR on embedded image {filename}: {e}]"
        
    return ""

def extract_text_from_excel_detailed(file_path: str, llm=None) -> dict:
    """Extracts text and returns details of cell reading vs embedded image OCR."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Excel file not found at {file_path}")
    
    cell_blocks = []
    cell_err = None
    
    # Layer 1: Cell text
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            cell_blocks.append(f"--- Sheet: {sheet_name} ---")
            for row in sheet.iter_rows(values_only=True):
                row_vals = [str(val).strip() for val in row if val is not None]
                if row_vals:
                    cell_blocks.append(" | ".join(row_vals))
    except Exception as e1:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=False)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                cell_blocks.append(f"--- Sheet (Formula Mode): {sheet_name} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(val).strip() for val in row if val is not None]
                    if row_vals:
                        cell_blocks.append(" | ".join(row_vals))
        except Exception as e2:
            cell_err = f"[Error reading spreadsheet cells: {e1} / {e2}]"
            cell_blocks.append(cell_err)
            
    text_cells = "\n".join(cell_blocks)
    
    # Layer 2: Images OCR
    image_blocks = []
    ocr_triggered = False
    images_found = []
    
    if zipfile.is_zipfile(file_path):
        try:
            with zipfile.ZipFile(file_path, 'r') as zip_ref:
                image_files = [f for f in zip_ref.namelist() if f.startswith('xl/media/')]
                images_found = [os.path.basename(img) for img in image_files]
                if image_files:
                    ocr_triggered = True
                    for img_name in sorted(image_files):
                        image_blocks.append(f"\n[Processing Embedded Image: {os.path.basename(img_name)}]")
                        try:
                            img_bytes = zip_ref.read(img_name)
                            if llm is not None:
                                ocr_txt = perform_image_ocr(img_bytes, img_name, llm)
                                if ocr_txt:
                                    image_blocks.append(ocr_txt)
                            else:
                                image_blocks.append("[AI OCR skipped: LLM provider not available]")
                        except Exception as img_err:
                            image_blocks.append(f"[Error processing image {img_name}: {img_err}]")
        except Exception as zip_err:
            image_blocks.append(f"[Error scanning spreadsheet archive for images: {zip_err}]")
            
    text_images = "\n".join(image_blocks)
    
    # Combine final text
    final_blocks = [text_cells]
    if text_images:
        final_blocks.append("\n--- Found Embedded Images in Spreadsheet ---" + text_images)
    final_merged = "\n".join(final_blocks)
    
    return {
        "text": final_merged.strip(),
        "details": {
            "layer_1_cells": text_cells,
            "ocr_triggered": ocr_triggered,
            "images_found": images_found,
            "layer_2_images_ocr": text_images,
            "final_merged": final_merged
        }
    }

def extract_text_from_excel(file_path: str, llm=None) -> str:
    """Extracts text from all sheets, cells, and embedded images of an Excel file."""
    res = extract_text_from_excel_detailed(file_path, llm=llm)
    return res["text"]
